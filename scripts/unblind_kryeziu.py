"""Explicit first outcome projection after protocol and implementation freeze."""
from datetime import datetime, timezone
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'artifacts/kryeziu_confirmation_v1'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    preflight = json.loads((OUT/'preflight.json').read_text())
    for name, expected in preflight['code_sha256'].items():
        if digest(ROOT/name) != expected:
            raise ValueError('Implementation changed after preflight: '+name)
    freeze_path = ROOT/'artifacts/kryeziu_design_v1/protocol-freeze.json'
    if digest(freeze_path) != preflight['protocol_freeze_sha256']:
        raise ValueError('Protocol freeze changed')
    for name, expected in json.loads(freeze_path.read_text())['sha256'].items():
        if digest(ROOT/name) != expected:
            raise ValueError('Frozen input changed: '+name)
    source = next(f for f in json.loads((ROOT/'data/kryeziu_candidate_metadata.json').read_text())['files']
                  if f['filename'] == 'Data S4.xlsx')
    raw = ROOT/'data/raw/kryeziu/Data S4.xlsx'
    if digest(raw) != source['sha256']:
        raise ValueError('Source workbook hash mismatch')
    started = {'unblinding_started_utc': datetime.now(timezone.utc).isoformat(),
               'preflight_sha256': digest(OUT/'preflight.json'), 'source_sha256': source['sha256']}
    with (OUT/'unblinding-started.json').open('x') as handle:
        json.dump(started, handle, indent=2)
    workbook = load_workbook(raw, read_only=True, data_only=True)
    path = ROOT/'data/raw/kryeziu/signals-projection.csv'
    count = 0
    try:
        if workbook.sheetnames != ['DSRT_RAW_211PDOs']:
            raise ValueError('Unexpected response worksheet structure')
        rows = workbook['DSRT_RAW_211PDOs'].iter_rows(values_only=True)
        header = next(rows)
        column = header.index('signal')
        with path.open('x', newline='') as handle:
            writer = csv.writer(handle)
            writer.writerow(['source_row', 'signal'])
            for source_row, row in enumerate(rows, 2):
                writer.writerow([source_row, row[column]])
                count += 1
    finally:
        workbook.close()
    result = {**started, 'sha256': digest(path), 'rows': count,
              'status': 'outcomes projected for first frozen-protocol evaluation; no tuning authorized'}
    (OUT/'signals-manifest.json').write_text(json.dumps(result, indent=2)+'\n')
    print(json.dumps(result, indent=2))


if __name__ == '__main__':
    main()
