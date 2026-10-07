"""Project public sample/patient identifiers without exposing mutation fields."""
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from inspect_kryeziu_design import acquire

ROOT = Path(__file__).resolve().parents[1]


def main():
    manifest = json.loads((ROOT/'data/kryeziu_candidate_metadata.json').read_text())
    source = next(f for f in manifest['files'] if f['filename'] == 'Data S1.xlsx')
    path = ROOT/'data/raw/kryeziu/Data S1.xlsx'
    acquire(source, path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    pairs = set()
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows)
            fields = ['sample_id', 'patient', 'sample_type']
            if not set(fields).issubset(header) or len(header) != len(set(header)):
                raise ValueError('Unrecognized sample-map schema; no data values exposed')
            indices = [header.index(field) for field in fields]
            for row in rows:
                values = tuple(str(row[i]) if row[i] is not None else '' for i in indices)
                if any(values):
                    if not all(values):
                        raise ValueError('Incomplete source identifier mapping')
                    pairs.add(values)
    finally:
        workbook.close()
    output = ROOT/'artifacts/kryeziu_design_v1/patient-map.csv'
    with output.open('w', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(fields)
        writer.writerows(sorted(pairs))
    record = {'source_sha256': source['sha256'], 'projection_sha256': hashlib.sha256(output.read_bytes()).hexdigest(),
              'distinct_sample_patient_type_rows': len(pairs), 'fields_exported': fields,
              'mutation_fields_exposed_or_analyzed': False}
    output.with_suffix('.json').write_text(json.dumps(record, indent=2)+'\n')
    print(json.dumps(record, indent=2))


if __name__ == '__main__':
    main()
